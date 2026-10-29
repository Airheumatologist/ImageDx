from pathlib import Path
import json,collections,re,datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
B=Path(__file__).parent;O=B.parent/'final_new_topics';O.mkdir(exist_ok=True)
adds=json.loads((B/'integrated_additions.json').read_text());dec=json.loads((B/'integrated_decisions.json').read_text());existing=json.loads((B/'existing_catalog.json').read_text());en={x['topic_id']:x['name'] for x in existing};A={x['topic_id']:x for x in adds}
# Keep only positively checked supporting references in the released package.
root_checked={u for x in json.loads((B/'review_neuro_rheum.json').read_text()) for u in x.get('source_urls',[])}|{'https://pubmed.ncbi.nlm.nih.gov/30080992/','https://academic.oup.com/eurheartj/article/44/37/3503/7246608'}
verified=root_checked|{u for x in dec if x.get('evidence_status')=='supported' for u in x.get('source_urls',[])}
for x in adds:
 x['source_urls']=[u for u in x['source_urls'] if u in verified]
 x['reference_status']='Selected supporting reference checked' if x['source_urls'] else 'Editorial recommendation without a separately verified topic reference'
 x['image_readiness']='Representative image selection, diagnosis confirmation, and reuse-rights clearance required before publication'
for i,x in enumerate(adds,1):x['list_number']=i
for x in dec:
 x['source_urls']=[u for u in x.get('source_urls',[]) if u in verified]
 if x['final_id'] in A:x['publication_outcome']='Covered by new topic '+A[x['final_id']]['name']
 elif x['action']=='MERGE_EXISTING':x['publication_outcome']='Expand existing topic '+', '.join(en[q] for q in x['existing_parent_ids'])
 elif x['action']=='MERGE_NEW':x['publication_outcome']='Include within new topic '+A.get(x.get('merge_target_id'),{}).get('name',x.get('merge_target_id',''))
 else:x['publication_outcome']='Do not add in this release'
# Canonical unique publication list and the decision record are portable deliverables.
counts=collections.Counter(x['priority'] for x in adds);sc=collections.defaultdict(collections.Counter)
for x in adds:sc[x['specialty']][x['priority']]+=1
metadata={'title':'Final new medical topics','review_date':'2026-10-28','baseline_catalog_topics':1643,'baseline_catalog_version':'4.1.0','scope':'Proposals in Newtopics.docx, reconciled against data/main_database_topics.json','total_additions':len(adds),'first_priority':counts['P1'],'next_priority':counts['P2'],'priority_basis':'Editorial clinical-learning and visual-teaching value, not a quantitative prevalence or evidence score','image_assets_verified':False,'topic_reference_coverage':'Selected targeted references; this is not an evidence review for every entry'}
(O/'final_new_topics.json').write_text(json.dumps({'metadata':metadata,'topics':adds},indent=2,ensure_ascii=False))
(O/'topic_decisions.json').write_text(json.dumps({'metadata':metadata,'decisions':dec},indent=2,ensure_ascii=False))
# Plain list deliberately has only addable canonical topics, without the deferred suggestions.
text=['FINAL NEW MEDICAL TOPICS',f'{len(adds)} additions | P1 {counts["P1"]} | P2 {counts["P2"]}','P1 = add first; P2 = add next.','']
for sp in sorted(sc):
 text+=[sp.upper()]
 for pr in ['P1','P2']:
  for x in adds:
   if x['specialty']==sp and x['priority']==pr:text.append(f"{x['list_number']:03}. [{pr}] {x['name']}")
 text.append('')
(O/'final_new_topics.txt').write_text('\n'.join(text))
# Full audit is searchable and keeps each source mention traceable.
def esc(s):return str(s).replace('|','/').replace('\n',' ')
md=['# Topic decision audit','',f"Final list: **{len(adds)} distinct additions** ({counts['P1']} P1; {counts['P2']} P2), compared with the 1,643-topic catalog dated October 28, 2026.",'','All 481 topic or group headings and the six final pulmonary bullet paragraphs were covered. Grouped paragraphs were expanded. The review produced '+str(len(dec))+' decision records; these are proposal mentions, not a count of unique new pages.','', 'The publication outcome is the final decision. Different source mentions can point to the same new page. A merge requires preserving the stated disease distinction and, where specified, broadening the existing title.','', 'References are selected checks, not an evidence dossier for every entry. Only references positively checked for topical relevance are retained here. A literature citation does not establish that its images can be reused.','', '## Inclusion rules','', '- A distinct disease, syndrome, injury, or structural diagnosis with a useful clinical or radiologic image.','- Clinical photographs, dermoscopy, fundus photographs, endoscopy, and radiology including ultrasound and echocardiography qualify.','- Histology, organism microscopy, laboratory tests, ECG, or genetic findings alone do not qualify.','- A typical appearance can be useful without being pathognomonic. The clinical context and confirmatory diagnosis remain essential.','- Keep subtypes and sites in one parent when they share the same visual teaching content. Retain clearly different disease entities when they add meaningful visual value.','', '## Catalog corrections and required title expansions','', '- Remove “angiosarcoma of bone” from the Adamantinoma synonym list. They are different tumors. Bone angiosarcoma is deferred for independent image review; the incorrect alias does not make it an existing topic.','- Broaden Hypertrophic Obstructive Cardiomyopathy to Hypertrophic Cardiomyopathy before adding apical nonobstructive images.','- Broaden Duchenne Muscular Dystrophy to Dystrophinopathies before adding Becker disease.','- Broaden Alpha-Thalassemia beyond the current HbH label before adding Hb Bart’s hydrops fetalis.','- Reorganize the MALT and ocular lymphoma content carefully: gastric MALT belongs to extranodal marginal zone lymphoma, while conjunctival lymphoma can include other histologies. Do not equate all conjunctival lymphomas with MALT.','- Broaden the existing pleural empyema necessitans topic to Pleural Empyema, with necessitans as a labeled complication.','- Broaden localized tracheobronchial amyloidosis to Localized Pulmonary Amyloidosis before adding nodular lung disease.','- Broaden Isthmic Spondylolisthesis to Spondylolisthesis before adding the degenerative subtype; preserve pars-defect and intact-pars distinctions.','- Preserve copper deficiency and nitrous oxide as different causes of dorsal-column myelopathy when expanding the existing subacute combined degeneration page.','', '## Decisions by source specialty','']
for sp in sorted({x['section'] for x in dec}):
 md+=['### '+sp,'','| Source proposal | Final outcome | Reason |','|---|---|---|']
 for x in dec:
  if x['section']==sp:md.append('| '+esc(x['canonical_name'])+' | '+esc(x['publication_outcome'])+' | '+esc(x['rationale'])+' |')
 md.append('')
(O/'topic_decision_audit.md').write_text('\n'.join(md))
# Word document with the complete addition list and clinically useful image anchors.
d=Document();sec=d.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
styles=d.styles
for sn in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=styles[sn];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 for border in list(s.element.xpath('.//w:pBdr')):border.getparent().remove(border)
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.line_spacing=1.05;styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Title'].font.size=Pt(25);styles['Title'].font.bold=True;styles['Title'].paragraph_format.space_after=Pt(10)
styles['Subtitle'].font.size=Pt(11);styles['Subtitle'].paragraph_format.space_after=Pt(10)
styles['Heading 1'].font.size=Pt(15);styles['Heading 1'].font.bold=True;styles['Heading 1'].paragraph_format.space_before=Pt(13);styles['Heading 1'].paragraph_format.space_after=Pt(6)
styles['Heading 2'].font.size=Pt(11.5);styles['Heading 2'].font.bold=True;styles['Heading 2'].paragraph_format.space_before=Pt(8);styles['Heading 2'].paragraph_format.space_after=Pt(5)
# Long reference list needs unobtrusive pagination.
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run();r.font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
d.add_paragraph('Final new medical topics','Title')
d.add_paragraph('Clinical and radiologic image atlas for physicians and medical students','Subtitle')
d.add_paragraph(f"Recommend {len(adds)} distinct additions: {counts['P1']} first-priority topics and {counts['P2']} next-priority topics. The list consolidates repeated proposals and assigns each topic one home specialty.")
d.add_paragraph('First priority covers core teaching diagnoses, common visual differentials, and important emergencies. Next priority extends specialty coverage with less common but useful clinical or imaging appearances. Both groups meet the editorial scope; priority reflects release order.')
d.add_paragraph('Selection criteria','Heading 1')
for s in ['Include distinct diseases, syndromes, injuries, and structural diagnoses with useful patient photographs, endoscopic or fundus images, dermoscopy, or radiologic manifestations. Ultrasound and echocardiography qualify.','Use pathology as supporting material. A topic whose useful image is only histology, organism microscopy, a laboratory result, or a genetic finding is deferred.','A teaching appearance need not be unique to one disease. Captions must state the relevant clinical context and avoid calling supportive findings diagnostic by themselves.','Keep repeated sites, causes, and overlapping phenotypes together when one page can teach them clearly. Record substantive additions to existing pages separately from new pages.']:
 d.add_paragraph(s,style='List Bullet')
d.add_paragraph('Before publication, select representative images, confirm the diagnosis and caption, and clear the required reuse rights. A listed topic or a cited article is not an approved image asset.')
d.add_paragraph('Addition counts by specialty','Heading 1')
t=d.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[3.4,.85,.85,.9]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for c,txt in zip(t.rows[0].cells,['Home specialty','First','Next','Total']):c.text=txt
for sp in sorted(sc):
 cells=t.add_row().cells
 for c,txt in zip(cells,[sp,str(sc[sp]['P1']),str(sc[sp]['P2']),str(sum(sc[sp].values()))]):c.text=txt
cells=t.add_row().cells
for c,txt in zip(cells,['Total',str(counts['P1']),str(counts['P2']),str(len(adds))]):c.text=txt
for ri,row in enumerate(t.rows):
 trpr=row._tr.get_or_add_trPr()
 if ri==0:
  rep=OxmlElement('w:tblHeader');trpr.append(rep)
 for ci,c in enumerate(row.cells):
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=c._tc.get_or_add_tcPr()
  b=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
  tcpr.append(b)
  margin=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margin.append(e)
  tcpr.append(margin)
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2' if ri==0 else ('F7F8FA' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
  for par in c.paragraphs:
   par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1
   if ci>0:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for run in par.runs:run.font.size=Pt(9);run.bold=(ri==0 or ri==len(t.rows)-1)
d.add_page_break()
d.add_paragraph('Topics to add','Heading 1')
d.add_paragraph('Each numbered entry is one new topic. The brief image anchor explains its clinical or radiologic teaching value. Related entities included within a shared topic are labeled in the anchor or decision audit.')
for sp in sorted(sc):
 d.add_paragraph(sp,'Heading 1')
 for pr,label in [('P1','First priority'),('P2','Next priority')]:
  items=[x for x in adds if x['specialty']==sp and x['priority']==pr]
  if not items:continue
  d.add_paragraph(f'{label} {len(items)} '+('topic' if len(items)==1 else 'topics'),'Heading 2')
  for x in items:
   par=d.add_paragraph();par.paragraph_format.keep_together=True;par.paragraph_format.space_after=Pt(5)
   par.add_run(f"{x['list_number']:03}  {x['name']}. ").bold=True
   par.add_run(x['visual_anchor'])
d.add_page_break();d.add_paragraph('Content for existing pages','Heading 1')
d.add_paragraph('These proposals add useful content to current topics. They are excluded from the new-topic count. Broaden a title where needed and keep different diseases or etiologies clearly labeled.')
by_parent=collections.defaultdict(list)
for x in dec:
 if x['action']=='MERGE_EXISTING':
  for parent in x['existing_parent_ids']:by_parent[parent].append(x)
for parent,items in sorted(by_parent.items(),key=lambda z:en[z[0]].lower()):
 names=list(dict.fromkeys(x['canonical_name'] for x in items));par=d.add_paragraph();par.paragraph_format.keep_together=True
 par.add_run(en[parent]+'. ').bold=True
 if parent=='intravascular_large_b_cell_lymphoma':par.add_run('Expand neurologic MRI coverage; the disease itself is already listed.')
 elif parent=='conjunctival_lymphoma':par.add_run('Reorganize gastric and ocular MALT content under extranodal marginal zone lymphoma; preserve other conjunctival lymphoma types.')
 elif parent=='urachal_carcinoma':par.add_run('Broaden title to Urachal Anomalies and Carcinoma before adding benign urachal cysts as a clearly separate section.')
 else:par.add_run('Add '+ '; '.join(names)+'.')
d.add_paragraph('Required catalog corrections','Heading 1')
for s in [
'Adamantinoma has an incorrect angiosarcoma-of-bone alias. Remove that alias; the two tumors are different. Bone angiosarcoma remains deferred for independent clinical and imaging review.',
'Use Hypertrophic Cardiomyopathy as the parent for obstructive and apical phenotypes. Use Dystrophinopathies for the Duchenne and Becker spectrum. Broaden Alpha-Thalassemia beyond the current HbH label before including Hb Bart\'s hydrops fetalis.',
'For MALT lymphoma, reorganize gastric and ocular marginal-zone disease together while preserving other conjunctival lymphoma histologies. Do not relabel every conjunctival lymphoma as MALT.',
'Broaden Pleural Empyema Necessitans to Pleural Empyema, and Localized Tracheobronchial Amyloidosis to Localized Pulmonary Amyloidosis, before adding their proposed related forms.',
'For degenerative spondylolisthesis, broaden the current isthmic title and preserve the distinction between a pars defect and an intact pars. For dorsal-column myelopathy, distinguish copper deficiency from functional vitamin B12 inactivation by nitrous oxide.',
'Use Alveolar Macrophage Pneumonia as the current ERS/ATS term, with Desquamative Interstitial Pneumonia and DIP retained for search. Treat isolated left ventricular hypertrabeculation as a morphology requiring clinical interpretation, rather than automatically labeling it a separate cardiomyopathy.'
]:d.add_paragraph(s,style='List Bullet')
d.add_paragraph('Selected clinical references','Heading 1')
d.add_paragraph('These references support selected imaging patterns, nomenclature, and consolidation decisions. They are reading references; image reuse requires its own license assessment.')
refs=[
('2025 ERS ATS statement on interstitial pneumonia classification','https://publications.ersnet.org/content/erj/66/6/2500158'),
('2023 ESC guidelines for cardiomyopathies','https://academic.oup.com/eurheartj/article/44/37/3503/7246608'),
('2022 ESC ERS guidelines for pulmonary hypertension','https://publications.ersnet.org/content/erj/61/1/2200879'),
('2023 EAN PNS guideline on Guillain Barre syndrome','https://onlinelibrary.wiley.com/doi/full/10.1111/ene.16073'),
('Clinical and radiologic features of Balo concentric sclerosis','https://pmc.ncbi.nlm.nih.gov/articles/PMC8362356/'),
('Enterovirus 71 infection and neurologic complications','https://pmc.ncbi.nlm.nih.gov/articles/PMC5099286/'),
('Dorsal thoracic arachnoid web and the scalpel sign','https://pmc.ncbi.nlm.nih.gov/articles/PMC7964642/'),
('GeneReviews on biotin thiamine responsive basal ganglia disease','https://www.ncbi.nlm.nih.gov/books/NBK169615/'),
('GeneReviews on HTRA1 disorder','https://www.ncbi.nlm.nih.gov/sites/books/NBK32533/'),
('GeneReviews on CSF1R related disorder','https://www.ncbi.nlm.nih.gov/books/NBK100239/'),
('GeneReviews on LAMA2 muscular dystrophy','https://www.ncbi.nlm.nih.gov/books/NBK97333/'),
('GeneReviews on multicentric osteolysis nodulosis and arthropathy','https://www.ncbi.nlm.nih.gov/books/NBK373578/'),
('Imaging findings of COPA syndrome','https://pubmed.ncbi.nlm.nih.gov/36746811/'),
('Clinical features of molecularly confirmed Sanjad Sakati syndrome','https://pubmed.ncbi.nlm.nih.gov/30080992/'),
]
def hyperlink(par,label,url):
 from docx.opc.constants import RELATIONSHIP_TYPE
 rid=par.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True);a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);a.append(r);par._p.append(a)
for i,(title,url) in enumerate(refs,1):par=d.add_paragraph();hyperlink(par,f'{i}. {title}',url)
d.core_properties.title='Final new medical topics';d.core_properties.subject='Clinical and radiologic atlas additions';d.core_properties.author='';d.core_properties.keywords='medical topics, clinical images, radiology, editorial review'
d.save(O/'final_new_topics.docx')
# Deterministic structural checks.
assert len(adds)==len({x['topic_id'] for x in adds})
assert all(x['visual_anchor'] and x['priority'] in ['P1','P2'] for x in adds)
assert set(x['topic_id'] for x in adds).isdisjoint(en)
assert all(x.get('merge_target_id') in A for x in dec if x['action']=='MERGE_NEW')
assert {x['uid'] for x in json.loads((B/'all_headings.json').read_text())}<={x.get('uid') for x in dec}
print(json.dumps(metadata));print('Written',[(p.name,p.stat().st_size) for p in O.iterdir() if p.is_file()])
